"""Build AlphaMaster beginner Word tutorial with embedded screenshots."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
IMG = ROOT / "docs" / "tutorial_images"
OUT = ROOT / "docs" / "AlphaMaster使用教學.docx"


def set_run_font(run, *, size=11, bold=False, color=None, name="微軟雅黑"):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    if color is not None:
        run.font.color.rgb = color
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), name)


def add_title(doc: Document, text: str):
    p = doc.add_heading(text, level=0)
    for run in p.runs:
        set_run_font(run, size=22, bold=True)


def add_h1(doc: Document, text: str):
    p = doc.add_heading(text, level=1)
    for run in p.runs:
        set_run_font(run, size=16, bold=True)


def add_h2(doc: Document, text: str):
    p = doc.add_heading(text, level=2)
    for run in p.runs:
        set_run_font(run, size=13, bold=True)


def add_p(doc: Document, text: str, *, bold=False, size=11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.35
    return p


def add_steps(doc: Document, items: list[str]):
    for i, item in enumerate(items, 1):
        p = doc.add_paragraph(style="List Number")
        run = p.add_run(item)
        set_run_font(run, size=11)
        p.paragraph_format.space_after = Pt(4)


def add_bullets(doc: Document, items: list[str]):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run_font(run, size=11)
        p.paragraph_format.space_after = Pt(3)


def add_tip(doc: Document, text: str):
    p = doc.add_paragraph()
    run = p.add_run("小建議：")
    set_run_font(run, size=11, bold=True, color=RGBColor(0x0B, 0x6E, 0x4F))
    run2 = p.add_run(text)
    set_run_font(run2, size=11, color=RGBColor(0x1F, 0x3A, 0x33))
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)


def add_warn(doc: Document, text: str):
    p = doc.add_paragraph()
    run = p.add_run("注意：")
    set_run_font(run, size=11, bold=True, color=RGBColor(0xB4, 0x3B, 0x2E))
    run2 = p.add_run(text)
    set_run_font(run2, size=11)
    p.paragraph_format.space_after = Pt(8)


def add_img(doc: Document, name: str, *, width_cm: float = 15.5, caption: str | None = None):
    path = IMG / name
    if not path.exists():
        add_p(doc, f"（缺少截圖：{name}）", bold=True)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Cm(width_cm))
    if caption:
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c.add_run(caption)
        set_run_font(r, size=9, color=RGBColor(0x55, 0x65, 0x70))
        c.paragraph_format.space_after = Pt(10)


def build() -> Path:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    add_title(doc, "AlphaMaster 超詳細使用教學（帶圖）")
    add_p(doc, "適合人群：第一次接觸本軟體的同學（盡量寫得像說明書）。", bold=True)
    add_p(
        doc,
        "這份教學會帶你走完完整流程：安裝與啟動 → 準備數據 → 訓練挖因子 → 回測檢驗 → 即時看信號。",
    )
    add_p(doc, f"對應軟體路徑範例：D:\\cl\\AlphaMaster　　網頁地址：http://127.0.0.1:8765")
    add_p(doc, "文件生成日期：以你電腦上打開的軟體界面為準（本教學插圖來自實際截圖）。")

    # 目錄式大綱
    add_h1(doc, "目錄")
    add_bullets(
        doc,
        [
            "一、這個軟體是幹什麼的？",
            "二、開始前要準備什麼",
            "三、怎麼安裝依賴（第一次必做）",
            "四、怎麼啟動軟體",
            "五、打開後頁面長什麼樣",
            "六、準備 K 線數據文件（最關鍵！）",
            "七、步驟① 訓練：讓電腦幫你挖因子",
            "八、步驟② 回測：用歷史數據檢驗策略",
            "九、步驟③ 即時分析：盯盤出信號",
            "十、常用按鈕一覽表",
            "十一、遇到報錯怎麼辦",
            "十二、建議的最佳實踐",
        ],
    )

    add_h1(doc, "一、這個軟體是幹什麼的？")
    add_p(
        doc,
        "一句話：AlphaMaster 會用人工智慧（強化學習）在歷史行情裡“挖”交易因子公式，再幫你回測和即時觀察信號。",
    )
    add_p(doc, "你可以把它想成下面三步：")
    add_steps(
        doc,
        [
            "訓練：餵給它一份 K 線數據，它拚命搜索更好的公式。",
            "回測：把找到的公式放到歷史行情上驗成績（看賺不賺錢、穩不穩）。",
            "即時分析：盯著當前市場，按同一套公式報方向（漲/跌/觀望）。",
        ],
    )
    add_tip(doc, "你不需要會寫程式碼。多數時間只需點按鈕、選文件、看圖表。")

    add_h1(doc, "二、開始前要準備什麼")
    add_bullets(
        doc,
        [
            "一台 Windows 電腦（本教學按 Windows 寫）。",
            "已安裝 Python 3.10 及以上（推薦 3.11 或 3.13）。",
            "AlphaMaster 項目文件夾（例如 D:\\cl\\AlphaMaster 或 D:\\Al_master）。",
            "至少一份行情 Parquet 文件（後面教你命名）。",
            "如果要用 MT5 即時數據：電腦上還要打開並登錄 MetaTrader 5。",
            "如果要用 AI 分析：準備 DeepSeek Key，或本機已開 QClaw / OpenClaw 一類本地助手。",
        ],
    )

    add_h1(doc, "三、怎麼安裝依賴（第一次必做）")
    add_p(doc, "打開「命令提示字元」或 PowerShell，進入項目文件夾，輸入：")
    add_p(doc, "cd /d D:\\cl\\AlphaMaster", bold=True)
    add_p(doc, "python -m pip install -r requirements.txt", bold=True)
    add_warn(
        doc,
        "請用 python -m pip，不要直接敲 pip（中文 Windows 上容易因編碼報錯）。",
    )
    add_p(doc, "裝完可用下面一行檢查關鍵包是否在：")
    add_p(
        doc,
        'python -c "import torch,fastapi,MetaTrader5,pyarrow,matplotlib,multipart; print(\'OK\')"',
        bold=True,
    )
    add_tip(
        doc,
        "TradingView 數據源若安裝失敗，可再執行："
        "python -m pip install git+https://github.com/rongardF/tvdatafeed.git",
    )

    add_h1(doc, "四、怎麼啟動軟體")
    add_h2(doc, "方法 A（推薦）：雙擊啟動腳本")
    add_steps(
        doc,
        [
            "進入項目文件夾，找到 start_web.bat。",
            "雙擊它。腳本會：先清理被占用的埠 → 啟動服務 → 等頁面真正就緒 → 自動打開瀏覽器。",
            "看到瀏覽器打開 http://127.0.0.1:8765 就成功了。",
            "啟動器窗口可以先留著；關掉啟動器窗口一般不會立刻關服務（服務在另一個“AlphaMaster Server”窗口）。",
        ],
    )
    add_h2(doc, "方法 B：命令行啟動")
    add_p(doc, "cd /d D:\\cl\\AlphaMaster", bold=True)
    add_p(doc, "python run_web.py --host 127.0.0.1 --port 8765", bold=True)
    add_p(doc, "然後在瀏覽器地址欄輸入：http://127.0.0.1:8765")
    add_warn(
        doc,
        "如果提示埠被占用，可以改成 --port 8766，或先關閉舊的 AlphaMaster Server 窗口。",
    )

    add_h1(doc, "五、打開後頁面長什麼樣")
    add_p(doc, "打開網頁後，最上方是軟體標題，中間有三個大步驟：01 模型訓練、02 策略回測、03 即時分析。")
    add_img(doc, "01_home_top.png", caption="圖1 打開軟體後的首頁上半部分")
    add_img(doc, "02_stepper.png", width_cm=14.5, caption="圖2 頂部三步導航：訓練 → 回測 → 即時")
    add_p(doc, "你要使用時，就從上到下按這三個步驟走。一般不要跳著亂點。")

    add_h1(doc, "六、準備 K 線數據文件（最關鍵！）")
    add_p(
        doc,
        "本軟體訓練和回測強制使用本地 Parquet 文件，不會在訓練時偷偷聯網拉行情。所以「檔案名正確」非常重要。",
    )
    add_h2(doc, "6.1 文件怎麼命名")
    add_p(doc, "標準格式：", bold=True)
    add_p(doc, "{品種}_{週期}.parquet", bold=True)
    add_p(doc, "正確例子：")
    add_bullets(
        doc,
        [
            "XAUUSD_H1.parquet（黃金 · 1小時）",
            "BTCUSDT_H1.parquet（比特幣 · 1小時）",
            "AAPL_H1.parquet（蘋果股票 · 1小時）",
            "002008_60min.parquet（A股代碼 · 60分鐘，軟體會自動認成 H1）",
            "600519_5min.parquet（會認成 M5）",
        ],
    )
    add_p(doc, "週期別名對照（你會用到的）：", bold=True)
    add_bullets(
        doc,
        [
            "H1 = 1小時 = 60min / 60m / 1h",
            "M5 = 5分鐘 = 5min / 5m",
            "M15 = 15分鐘 = 15min / 15m",
            "H4 = 4小時 = 4h / 240min",
            "D1 = 日線 = 1d / day / daily",
        ],
    )
    add_h2(doc, "6.2 文件裡面要有哪些列")
    add_p(doc, "Parquet 至少要有：time、open、high、low、close，以及 volume 或 tick_volume。")
    add_p(doc, "K 線數量也不能太少，否則會提示「數據不足」。")
    add_warn(doc, "不要選 Excel、CSV 直接訓練。必須是 .parquet。")

    add_h1(doc, "七、步驟① 訓練：讓電腦幫你挖因子")
    add_p(doc, "確保頂部導航停在「01 模型訓練」。")

    add_h2(doc, "7.1 選擇數據文件")
    add_img(doc, "03_train_launch.png", caption="圖3 訓練啟動區：選擇數據文件 + 各類按鈕")
    add_steps(
        doc,
        [
            "點擊「選擇數據文件」。",
            "在彈出的文件窗口裡，找到你的 xxx_H1.parquet（或 002008_60min.parquet）。",
            "選中後，中間卡片會顯示品種、週期、K線數量。如果寫「文件不存在」或紅色報錯，先檢查路徑和命名。",
            "只有選對文件後，「開始訓練」才會變成可點。",
        ],
    )

    add_h2(doc, "7.2 開始訓練")
    add_steps(
        doc,
        [
            "點擊藍色「開始訓練」。右上角狀態會從「空閒」變成訓練中。",
            "左側會出現「訓練曲線」：綠色是最優分數，藍色是驗證分數。",
            "右側「訓練日誌」會不斷刷出進度。",
            "下方「最優公式」會顯示目前找到的最好因子寫法。",
        ],
    )
    add_img(doc, "04_train_chart.png", caption="圖4 訓練曲線區域（開始訓練後會動）")
    add_img(doc, "05_train_log.png", caption="圖5 訓練日誌區域")
    add_img(doc, "06_train_full.png", width_cm=14.0, caption="圖6 訓練頁整頁示意")

    add_h2(doc, "7.3 兩個分數分別是什麼意思（很重要）")
    add_bullets(
        doc,
        [
            "最優分數：到目前為止挖到的「最好那一條公式」的成績。通常只升不降，抬一截說明挖到更好的了。",
            "驗證分數：最近一批公式在驗證集上的平均表現，會上下波動，屬正常。",
            "不要只因為驗證分數某一刻掉了就慌——軟體有時會主動「重新攪一攪」去探索新公式。",
        ],
    )

    add_h2(doc, "7.4 「開始訓練」和「重新訓練」有什麼區別")
    add_bullets(
        doc,
        [
            "開始訓練：接著上次的檢查點繼續挖（省時間）。",
            "重新訓練：清空檢查點，從零重新搜一輪（可能挖到更好的，但更花時間）。",
            "已有更好策略不會被隨便沖掉：只有新公式更強才會覆蓋。",
        ],
    )

    add_h2(doc, "7.5 什麼時候可以停")
    add_bullets(
        doc,
        [
            "最優分數長時間幾乎不動，可以考慮停止。",
            "點紅色「停止」。停止後，最好策略一般已保存在 strategies 文件夾。",
            "也可以繼續掛著挖，時間越長通常機會越多（但費電費時間）。",
        ],
    )

    add_h2(doc, "7.6 導出策略 / 導出訓練 / 導入訓練")
    add_bullets(
        doc,
        [
            "導出策略：下載當前品種最優策略 JSON，方便備份或給別人回測。",
            "導出訓練：把檢查點、曲線、策略打成 zip，換電腦可繼續。",
            "導入訓練：上傳以前導出的 zip 或 .pt，下次可斷點續訓。",
        ],
    )

    add_h2(doc, "7.7 已保存策略列表")
    add_img(doc, "07_strategies.png", caption="圖7 已保存策略表：能看到品種、週期、分數、公式")
    add_p(doc, "訓練成功產生的策略會出現在這裡。回測和即時分析經常會用到這些文件。")

    add_h2(doc, "7.8 AI 分析（可選）")
    add_img(doc, "08_ai_panel.png", caption="圖8 AI 分析面板")
    add_steps(
        doc,
        [
            "若使用 DeepSeek：在輸入框填入 API Key。",
            "若本機有 QClaw / OpenClaw：可按界面提示切換 provider。",
            "點擊「開始分析」，AI 會用白話解釋：現在訓練情況如何、值不值得繼續、因子大概在幹什麼。",
        ],
    )
    add_tip(doc, "AI 分析是助手，不是買賣建議。最終仍要用回測結果自己判斷。")

    add_h1(doc, "八、步驟② 回測：用歷史數據檢驗策略")
    add_p(doc, "點頂部「02 策略回測」。")
    add_img(doc, "09_backtest_top.png", caption="圖9 回測頁頂部")
    add_img(doc, "10_backtest_launch.png", caption="圖10 選擇策略、設置手續費/滑點、開始回測")

    add_h2(doc, "8.1 怎麼做回測")
    add_steps(
        doc,
        [
            "點擊「選擇策略」，選 strategies 裡的 best_品種.json，或你導出的策略 JSON。",
            "（可選）改手續費、滑點。默認單邊手續費 0.02%、滑點 0.01%，合計約 0.03%。",
            "確認該策略能找到對應 Parquet（策略裡常記錄 data_file；沒有則回退訓練頁同品種文件）。",
            "點「開始回測」，等日誌跑完。",
            "查看「回測績效」「績效明細」和下方資金曲線。",
        ],
    )
    add_img(doc, "11_backtest_summary.png", caption="圖11 回測績效區域（跑完後會出數字）")
    add_img(doc, "12_backtest_full.png", width_cm=14.0, caption="圖12 回測頁整頁示意")

    add_h2(doc, "8.2 回測結果怎麼讀（小學生版）")
    add_bullets(
        doc,
        [
            "收益：這段歷史上是賺還是虧。",
            "Sharpe / Sortino：穩不穩；越高一般越好（但也要結合交易次數看）。",
            "盈虧比：賺的時候平均賺多少，對比虧的時候平均虧多少。",
            "勝率：猜對方向的比例。勝率高不一定就好，還要看盈虧比。",
            "交易數：太少不夠信；多了也不等於一定穩。",
        ],
    )
    add_warn(
        doc,
        "回測好看 ≠ 未來一定賺錢。至少換一段數據、或換手續費成本再驗一次更踏實。",
    )
    add_warn(
        doc,
        "若報錯 No module named 'matplotlib'：執行 python -m pip install matplotlib 後重試。",
    )

    add_h1(doc, "九、步驟③ 即時分析：盯盤出信號")
    add_p(doc, "點頂部「03 即時分析」。")
    add_img(doc, "13_realtime_top.png", caption="圖13 即時分析頁頂部")
    add_img(doc, "14_realtime_form.png", caption="圖14 添加監控：數據源 + 品種 + 週期 + 策略因子")
    add_img(doc, "15_realtime_full.png", width_cm=14.0, caption="圖15 即時分析整頁示意")

    add_h2(doc, "9.1 添加一個監控怎麼操作")
    add_steps(
        doc,
        [
            "選擇數據源：常見是 MT5 或 TradingView（界面當前主要展示這兩者）。",
            "填寫品種，例如 XAUUSD、EURUSD。也可從下拉預設裡選。",
            "選擇週期，盡量和策略訓練時的週期一致（例如策略是 H1，這裡也選 1h）。",
            "選擇策略因子：選本機 best_xxx.json，或點導入策略。",
            "點擊添加。頁面會出現一張「信號卡片」。",
        ],
    )

    add_h2(doc, "9.2 信號卡片怎麼看")
    add_bullets(
        doc,
        [
            "預期上漲 / 預期下跌 / 先觀望：當前公式給出的方向。",
            "把握大小：強度翻譯成白話，告訴你這撥信號有多堅定。",
            "距離下次判斷：還有多久重新算一次。",
            "若狀態是「錯誤」：看卡片上的紅字，通常是數據源連不上或歷史不夠。",
        ],
    )

    add_h2(doc, "9.3 TradingView 連不上怎麼辦")
    add_p(
        doc,
        "如果添加 TradingView 監控時彈出「無法使用 TradingView」，按彈出視窗提示做：",
    )
    add_bullets(
        doc,
        [
            "把 VPN 設成全局，並打開 TUN（虛擬網卡）模式再試。",
            "或按彈出視窗「使用雲端伺服器」去看部署說明。",
            "或改回 MT5 數據源（本機 MT5 已登錄時更穩）。",
        ],
    )

    add_h2(doc, "9.4 飛書提醒（可選）")
    add_p(
        doc,
        "即時頁可配置飛書機器人：當方向發生轉折時推送文本提醒。按頁面上的幫助去開通 webhook 即可。",
    )

    add_h1(doc, "十、常用按鈕一覽表")
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "按鈕"
    hdr[1].text = "在哪一頁"
    hdr[2].text = "幹什麼"
    rows = [
        ("選擇數據文件", "訓練", "挑選本地 Parquet"),
        ("開始訓練", "訓練", "繼續挖因子"),
        ("重新訓練", "訓練", "清空檢查點重搜"),
        ("停止", "訓練/回測", "停掉當前任務"),
        ("導出策略", "訓練", "下載最優策略 JSON"),
        ("導出/導入訓練", "訓練", "遷移訓練進度"),
        ("開始分析", "訓練·AI區", "讓 AI 解釋訓練情況"),
        ("選擇策略", "回測", "挑選要檢驗的策略"),
        ("開始回測", "回測", "跑歷史績效"),
        ("添加監控", "即時", "掛一個即時信號任務"),
    ]
    for a, b, c in rows:
        cells = table.add_row().cells
        cells[0].text = a
        cells[1].text = b
        cells[2].text = c
    doc.add_paragraph()

    add_h1(doc, "十一、遇到報錯怎麼辦")
    add_h2(doc, "11.1 頁面一打開就一堆「網路錯誤 Failed to fetch」")
    add_bullets(
        doc,
        [
            "通常是服務剛啟動，還在載入 torch，還沒準備好。",
            "請用 start_web.bat 啟動（它會等介面就緒再開瀏覽器）。",
            "新版本前端會自動重試幾次；仍失敗就等 10 秒刷新頁面。",
        ],
    )
    add_h2(doc, "11.2 檔案名報錯（例如 002008_60min.parquet）")
    add_p(
        doc,
        "新版本已支持 60min/5min/1h 等別名。若仍報錯，請確認你用的是已更新的 parquet_manager，並已重啟 Web。",
    )
    add_h2(doc, "11.3 回測說沒有 matplotlib")
    add_p(doc, "python -m pip install matplotlib", bold=True)
    add_h2(doc, "11.4 選擇數據文件提示缺 pyarrow")
    add_p(doc, "python -m pip install pyarrow", bold=True)
    add_h2(doc, "11.5 歷史訓練時長一上來就很大")
    add_p(
        doc,
        "可能是倉庫自帶作者訓練記錄。清空 training_time_品種.json，並刪除 logs/train_品種_*.log 後再刷新。",
    )
    add_h2(doc, "11.6 MT5 數據源不可用")
    add_bullets(
        doc,
        [
            "確認已安裝 MetaTrader5 Python 包。",
            "確認 MT5 終端已打開並登入帳號。",
        ],
    )

    add_h1(doc, "十二、建議的最佳實踐（照著做不容易踩坑）")
    add_steps(
        doc,
        [
            "一次只訓練一個品種、一個週期（別同時開好多任務搶 CPU）。",
            "檔案名寫清楚：品種_週期.parquet。",
            "訓練時多看「最優分數」是否還在抬升。",
            "訓完先回測，再考慮即時。",
            "即時監控的週期盡量與策略訓練週期一致。",
            "定期導出策略和訓練包做備份。",
            "換電腦部署：複製項目 → python -m pip install -r requirements.txt → start_web.bat。",
        ],
    )

    add_h1(doc, "附錄：最短上手路線（10 分鐘版）")
    add_steps(
        doc,
        [
            "雙擊 start_web.bat，等瀏覽器打開。",
            "準備好 XAUUSD_H1.parquet（或你自己的數據）。",
            "訓練頁 → 選擇數據文件 → 開始訓練。",
            "過一陣點停止（或等分數抬升後停）。",
            "回測頁 → 選擇剛生成的策略 → 開始回測。",
            "即時頁 → 選 MT5 → 填同一品種與週期 → 選同一策略 → 添加。",
        ],
    )
    add_p(doc, "到這裡，你已經走完 AlphaMaster 的完整閉環。", bold=True)
    add_p(
        doc,
        "插圖目錄：docs/tutorial_images/　　本文件：docs/AlphaMaster使用教學.docx",
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    return OUT


if __name__ == "__main__":
    path = build()
    print("saved", path)
